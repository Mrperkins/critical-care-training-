"""Clinical review workbook from review/clinical-review-packet.json (written by app/scripts/review-packet.ts).

    python tools/review_packet_xlsx.py   →   review/clinical-review-packet.xlsx

Sheets: How to review (legend + one example row), Review (every item, Verdict dropdown + Comment, yellow = fill in),
Summary (live COUNTIFS per section).
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
rows = json.loads((ROOT / 'review/clinical-review-packet.json').read_text())
F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3A4D')
B = Font(name=F, size=10); BB = Font(name=F, size=10, bold=True)
FILL = PatternFill('solid', fgColor='FFFF00'); WRAP = Alignment(wrap_text=True, vertical='top')
THIN = Border(bottom=Side(style='thin', color='D0D7DD'))
VERDICTS = ['OK', 'Change', 'Remove']

wb = Workbook()
# ── How to review ──────────────────────────────────────────────────────────────────────────────────────
g = wb.active; g.title = 'How to review'
legend = [
    ('Clinical review — Critical Care Physiology', ''),
    ('', ''),
    ('What this is', 'Every teaching claim in the app, one row each: reference values, all lesson narration, challenge questions with their answers and explanations, real clinical media captions and overlays, and the descriptions shown when a 3D label is tapped.'),
    ('What to fill in', 'Only the two yellow columns on the Review sheet: Verdict (pick OK, Change or Remove from the list) and Reviewer comment (what is wrong and what it should say).'),
    ('Where to look', '"Where in the app" names the module and screen. Overlay rows ask you to check placement on the image in the app itself.'),
    ('Progress', 'The Summary sheet counts verdicts per section as you go.'),
    ('Source', f'Generated from the app source by app/scripts/review-packet.ts ({len(rows)} rows). Regenerated on every site build; the IDs stay stable so comments can be matched back.'),
    ('', ''),
    ('Example row (not part of the review)', ''),
]
for r, (a, b) in enumerate(legend, 1):
    g.cell(r, 1, a).font = Font(name=F, size=14, bold=True) if r == 1 else BB
    g.cell(r, 2, b).font = B; g.cell(r, 2).alignment = WRAP
ex_h = ['Section', 'ID', 'Where in the app', 'Item', 'Content to review', 'Verdict', 'Reviewer comment']
ex = ['1 Reference values', 'vent:PEEP', 'Ventilator › Normal values › Settings', 'PEEP', '5 cmH₂O — Physiological minimum…', 'Change', 'Say "typical starting PEEP" rather than "physiological minimum".']
for c, (h, v) in enumerate(zip(ex_h, ex), 1):
    g.cell(10, c, h).font = BB; g.cell(11, c, v).font = Font(name=F, size=10, italic=True, color='666666'); g.cell(11, c).alignment = WRAP
g.column_dimensions['A'].width = 34; g.column_dimensions['B'].width = 100
for c in 'CDEFG': g.column_dimensions[c].width = 22

# ── Review ─────────────────────────────────────────────────────────────────────────────────────────────
ws = wb.create_sheet('Review')
head = ['Section', 'ID', 'Where in the app', 'Item', 'Content to review', 'Source / credit', 'Verdict', 'Reviewer comment']
widths = [18, 26, 34, 34, 80, 34, 12, 44]
for c, (h, w) in enumerate(zip(head, widths), 1):
    cell = ws.cell(1, c, h); cell.font = H; cell.fill = HF; cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.column_dimensions[get_column_letter(c)].width = w
for i, r in enumerate(rows, 2):
    vals = [r['section'], r['id'], r['where'], r['item'], r['content'], r.get('source', ''), None, None]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(i, c, v); cell.font = B; cell.alignment = WRAP; cell.border = THIN
    ws.cell(i, 7).fill = FILL; ws.cell(i, 8).fill = FILL
n = len(rows) + 1
dv = DataValidation(type='list', formula1='"' + ','.join(VERDICTS) + '"', allow_blank=True, showErrorMessage=True, errorTitle='Verdict', error='Choose OK, Change or Remove.')
ws.add_data_validation(dv); dv.add(f'G2:G{n}')
ws.freeze_panes = 'C2'; ws.auto_filter.ref = f'A1:H{n}'; ws.row_dimensions[1].height = 30

# ── Summary (formulas) ─────────────────────────────────────────────────────────────────────────────────
sm = wb.create_sheet('Summary')
sec = sorted({r['section'] for r in rows})
hdr = ['Section', 'Items', 'OK', 'Change', 'Remove', 'Not yet reviewed', '% reviewed']
for c, h in enumerate(hdr, 1):
    cell = sm.cell(1, c, h); cell.font = H; cell.fill = HF
for i, s in enumerate(sec, 2):
    sm.cell(i, 1, s).font = B
    sm.cell(i, 2, f'=COUNTIF(Review!$A$2:$A${n},$A{i})')
    for c, v in zip((3, 4, 5), VERDICTS):
        sm.cell(i, c, f'=COUNTIFS(Review!$A$2:$A${n},$A{i},Review!$G$2:$G${n},"{v}")')
    sm.cell(i, 6, f'=B{i}-C{i}-D{i}-E{i}')
    sm.cell(i, 7, f'=IF(B{i}=0,0,(C{i}+D{i}+E{i})/B{i})').number_format = '0.0%'
    for c in range(2, 8): sm.cell(i, c).font = B
t = len(sec) + 2
sm.cell(t, 1, 'Total').font = BB
for c in range(2, 7):
    L = get_column_letter(c); sm.cell(t, c, f'=SUM({L}2:{L}{t - 1})').font = BB
sm.cell(t, 7, f'=IF(B{t}=0,0,(C{t}+D{t}+E{t})/B{t})').number_format = '0.0%'; sm.cell(t, 7).font = BB
sm.column_dimensions['A'].width = 26
for c in 'BCDEFG': sm.column_dimensions[c].width = 16
wb.move_sheet('Summary', offset=-1)

out = ROOT / 'review/clinical-review-packet.xlsx'; wb.save(out); print('wrote', out, len(rows), 'rows')
